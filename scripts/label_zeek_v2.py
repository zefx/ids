#!/usr/bin/env python3
"""
label_zeek_v2.py — Label Zeek conn.log using CICIDS2018 attack schedule Excel.

Key fixes vs label_zeek.py:
  - AST timezone (UTC-4, Atlantic Standard Time = UNB/Fredericton timezone)
  - Interval intersection: flow_start <= attack_end AND flow_end >= attack_start
  - ALL IPs from attacker/victim columns (both 172.31.x.x internal and public)
  - Bidirectional IP check (needed for Bot: victims→C2)

Usage:
  python3 label_zeek_v2.py --conn /data/cicids_zeek/Friday-16-02-2018/conn.log \
                           --schedule cse_cic_ids2018_attack_schedule.xlsx \
                           --day 2018-02-16 --out labeled_feb16.csv

Dependencies:
  pip install openpyxl
"""

import argparse
import csv
import re
import sys
from datetime import datetime, timedelta, timezone
from collections import Counter

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

# Atlantic Standard Time (UTC-4) — UNB Fredericton, New Brunswick
AST = timezone(timedelta(hours=-4))

# Fields to keep from conn.log (plus 'label' appended)
KEEP_FIELDS = [
    "ts", "uid", "id.orig_h", "id.orig_p", "id.resp_h", "id.resp_p",
    "proto", "service", "duration", "orig_bytes", "resp_bytes", "conn_state",
    "orig_pkts", "resp_pkts", "orig_ip_bytes", "resp_ip_bytes", "history"
]

# Application-layer attacks excluded from labeling:
# conn.log captures TCP connections, not HTTP requests. HTTP keep-alive means
# thousands of L7 requests = 1 conn.log row → too few samples for training.
# These attacks require http.log (future work).
EXCLUDED_ATTACKS = {
    "Brute Force -Web",
    "Brute Force -XSS",
    "SQL Injection",
    "Infiltration",
    "DDOS-LOIC-UDP",   # 120 flows total — too few for ML (UDP flood nature)
}

_IP_RE = re.compile(r'\b(\d{1,3}(?:\.\d{1,3}){3})\b')


def all_ips(raw):
    """Extract ALL IPs from a string — both internal (172.31.x.x) and public.
    Traffic within the AWS VPC uses internal IPs in pcap; external attackers
    (no 172.31 prefix in Excel) use public IPs directly.
    Including all IPs covers both cases.
    """
    return set(m.group(1) for m in _IP_RE.finditer(str(raw)))


def load_windows(xlsx_path, filter_date=None):
    """
    Load attack windows from Excel.
    Expected sheet name: 'Attack Schedule'
    Expected columns: Attack_Start, Attack_Finish, Attack_Name, Attacker_Raw, Victim_Raw

    filter_date: ISO string 'YYYY-MM-DD' — only load windows for this day (AST date).
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # Find the sheet — try exact name first, then first sheet
    if "Attack Schedule" in wb.sheetnames:
        ws = wb["Attack Schedule"]
    else:
        ws = wb.worksheets[0]
        print(f"[warn] Sheet 'Attack Schedule' not found, using '{ws.title}'", file=sys.stderr)

    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers) if h is not None}

    required = {"Attack_Start", "Attack_Finish", "Attack_Name", "Attacker_Raw", "Victim_Raw"}
    missing = required - set(col.keys())
    if missing:
        sys.exit(f"Missing columns in Excel: {missing}\nFound: {list(col.keys())}")

    windows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        attack_start = row[col["Attack_Start"]]
        attack_finish = row[col["Attack_Finish"]]
        attack_name = str(row[col["Attack_Name"]] or "").strip()
        attacker_raw = str(row[col["Attacker_Raw"]] or "")
        victim_raw = str(row[col["Victim_Raw"]] or "")

        if not isinstance(attack_start, datetime) or not isinstance(attack_finish, datetime):
            continue
        if not attack_name:
            continue

        # Treat Excel datetime as AST (no tzinfo in Excel — we assign it)
        start_ast = attack_start.replace(tzinfo=AST)
        end_ast = attack_finish.replace(tzinfo=AST)

        if filter_date and start_ast.date().isoformat() != filter_date:
            continue

        if attack_name in EXCLUDED_ATTACKS:
            print(f"      [excluded] {attack_name} ({start_ast.date()})", file=sys.stderr)
            continue

        attackers = all_ips(attacker_raw)
        victims = all_ips(victim_raw)

        windows.append({
            "label": attack_name,
            "attackers": attackers,
            "victims": victims,
            "start": start_ast,
            "end": end_ast,
        })

    return windows


def get_label(ts_epoch, duration_str, src_ip, dst_ip, windows):
    """
    Assign label to a flow.
    Uses interval intersection: flow overlaps window if flow_start <= w.end AND flow_end >= w.start.
    Checks both forward (attacker→victim) and reverse (victim→attacker) directions.
    """
    flow_start = datetime.fromtimestamp(ts_epoch, tz=AST)

    try:
        dur = float(duration_str) if duration_str and duration_str not in ("-", "") else 0.0
    except ValueError:
        dur = 0.0

    flow_end = flow_start + timedelta(seconds=dur) if dur > 0 else flow_start

    for w in windows:
        # Interval intersection
        if flow_start <= w["end"] and flow_end >= w["start"]:
            fwd = (src_ip in w["attackers"] and dst_ip in w["victims"])
            rev = (dst_ip in w["attackers"] and src_ip in w["victims"])
            if fwd or rev:
                return w["label"]

    return "BENIGN"


def process_conn_log(conn_path, windows, out_path):
    """Read Zeek conn.log (TSV with #fields header), write labeled CSV."""
    fields = None
    stats = Counter()
    by_label = Counter()

    with open(conn_path, "r") as fin, open(out_path, "w", newline="") as fout:
        writer = None
        out_fields = None

        for line in fin:
            line = line.rstrip("\n\r")
            if not line:
                continue
            if line.startswith("#close") or line.startswith("#open"):
                continue
            if line.startswith("#fields"):
                fields = line.split("\t")[1:]
                continue
            if line.startswith("#") or fields is None:
                continue

            parts = line.split("\t")
            if len(parts) != len(fields):
                stats["skipped_malformed"] += 1
                continue

            row = dict(zip(fields, parts))

            if writer is None:
                out_fields = [f for f in KEEP_FIELDS if f in row] + ["label"]
                writer = csv.DictWriter(fout, fieldnames=out_fields)
                writer.writeheader()

            try:
                ts_epoch = float(row.get("ts", 0))
            except ValueError:
                stats["skipped_bad_ts"] += 1
                continue

            label = get_label(
                ts_epoch,
                row.get("duration", "-"),
                row.get("id.orig_h", ""),
                row.get("id.resp_h", ""),
                windows
            )

            stats["total"] += 1
            by_label[label] += 1

            out_row = {f: row.get(f, "-") for f in out_fields[:-1]}
            out_row["label"] = label
            writer.writerow(out_row)

    return stats, by_label


def main():
    parser = argparse.ArgumentParser(description="Label Zeek conn.log from CICIDS2018 attack schedule")
    parser.add_argument("--conn", required=True, help="Path to Zeek conn.log")
    parser.add_argument("--schedule", required=True, help="Path to attack schedule Excel (.xlsx)")
    parser.add_argument("--day", required=False, default=None,
                        help="Filter to single day YYYY-MM-DD (AST). Omit to load all windows.")
    parser.add_argument("--out", required=True, help="Output CSV path")
    args = parser.parse_args()

    print(f"[1/3] Loading attack windows from {args.schedule} ...", file=sys.stderr)
    windows = load_windows(args.schedule, filter_date=args.day)

    if not windows:
        print(f"[warn] No attack windows found for day={args.day}. All flows will be BENIGN.", file=sys.stderr)
    else:
        print(f"      Loaded {len(windows)} windows:", file=sys.stderr)
        for w in windows:
            print(f"        {w['label']:40s} {w['start'].strftime('%H:%M')}–{w['end'].strftime('%H:%M')} AST  "
                  f"attackers={w['attackers']}  victims_count={len(w['victims'])}", file=sys.stderr)

    print(f"[2/3] Processing {args.conn} ...", file=sys.stderr)
    stats, by_label = process_conn_log(args.conn, windows, args.out)

    print(f"[3/3] Done. Written to {args.out}", file=sys.stderr)
    print(f"\n=== Results ===", file=sys.stderr)
    print(f"Total flows:   {stats['total']:>10,}", file=sys.stderr)
    if stats["skipped_malformed"]:
        print(f"Skipped (malformed): {stats['skipped_malformed']:>10,}", file=sys.stderr)
    print(f"\nLabel distribution:", file=sys.stderr)
    for label, count in sorted(by_label.items(), key=lambda x: -x[1]):
        pct = count / stats["total"] * 100 if stats["total"] else 0
        print(f"  {label:45s} {count:>10,}  ({pct:.1f}%)", file=sys.stderr)


if __name__ == "__main__":
    main()
